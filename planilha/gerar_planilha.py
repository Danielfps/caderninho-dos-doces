"""Gera planilha/caderninho-planilha.xlsx (custos, preços e fiado).

Uso: python3 planilha/gerar_planilha.py
Precisa só do openpyxl. O arquivo sai com as fórmulas sem resultado guardado;
Google Planilhas e Excel calculam ao abrir. Para os números aparecerem também em
visualizadores simples (ex.: abrir pelo WhatsApp), recalcule com o LibreOffice:
  soffice --headless --convert-to xlsx --outdir /tmp planilha/caderninho-planilha.xlsx
e copie o resultado por cima. A estrutura das abas está em docs/planilha.md.
"""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "caderninho-planilha.xlsx")

# Cores do app (app/index.html, tema claro)
ROSA = "D6456F"          # --accent
ROSA_CLARO = "FBE0E8"    # --accent-soft  -> células que a pessoa preenche
ROSA_FUNDO = "F8E9EC"    # --surface-2
CHOCOLATE = "5A3426"     # --cocoa
TINTA = "3B2420"         # --ink
APAGADO = "7A5F59"       # --muted
LINHA = "ECD6DA"         # --line
VERDE = "2B7A4B"         # --good
VERDE_CLARO = "E0F2E6"   # --good-soft
LARANJA = "BF3F0A"       # --bad
LARANJA_CLARO = "FDE6D8" # --bad-soft
BRANCO = "FFFFFF"

FONTE = "Arial"
DINHEIRO = '"R$" #,##0.00'
PCT = "0%"
DATA = "dd/mm/yyyy"

# Quantas linhas cada tabela aceita
N_ING, N_REC, N_PROD, N_CLI, N_VEN, N_PAG = 200, 500, 50, 200, 2000, 2000
L0 = 5  # primeira linha de dados das tabelas (linha 4 = cabeçalho)

fino = Side(style="thin", color=LINHA)
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)


def fill(cor):
    return PatternFill("solid", start_color=cor, end_color=cor)


def font(cor=TINTA, size=11, bold=False, italic=False):
    return Font(name=FONTE, color=cor, size=size, bold=bold, italic=italic)


def topo(ws, titulo, explica, ate_col, cor_aba):
    """Faixa rosa com o nome da aba e uma frase explicando para que serve."""
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = cor_aba
    ws.column_dimensions["A"].width = 2.5
    ult = get_column_letter(ate_col)
    ws.merge_cells(f"B1:{ult}1")
    ws.merge_cells(f"B2:{ult}2")
    c = ws["B1"]
    c.value = titulo
    c.font = Font(name=FONTE, size=20, bold=True, color=BRANCO)
    c.alignment = Alignment(vertical="center", indent=1)
    c = ws["B2"]
    c.value = explica
    c.font = font(APAGADO, 11, italic=True)
    c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    for col in range(2, ate_col + 1):
        ws.cell(1, col).fill = fill(ROSA)
        ws.cell(2, col).fill = fill(ROSA_FUNDO)
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 10


def cabecalho(ws, linha, colunas):
    """colunas: lista de (titulo, largura, eh_entrada)."""
    for i, (titulo, larg, entrada) in enumerate(colunas, start=2):
        c = ws.cell(linha, i, titulo)
        c.font = Font(name=FONTE, size=11, bold=True, color=BRANCO)
        c.fill = fill(ROSA if entrada else CHOCOLATE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[linha].height = 44


def celula(ws, ref, valor=None, entrada=False, fmt=None, bold=False, cor=TINTA, alinh=None, fundo=None, size=11):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.font = font(cor, size, bold=bold)
    c.fill = fill(fundo or (ROSA_CLARO if entrada else BRANCO))
    c.border = BORDA
    c.alignment = Alignment(vertical="center", horizontal=alinh, wrap_text=True)
    if fmt:
        c.number_format = fmt
    return c


def grade(ws, l_ini, n, colunas_fmt):
    """Formata n linhas vazias da tabela. colunas_fmt: {letra: (entrada, formato, alinhamento)}."""
    for r in range(l_ini, l_ini + n):
        ws.row_dimensions[r].height = 21
        for col, (entrada, fmt, alinh) in colunas_fmt.items():
            celula(ws, f"{col}{r}", entrada=entrada, fmt=fmt, alinh=alinh)


def lista(ws, formula, intervalo, msg=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = False  # deixa escrever algo novo sem travar
    if msg:
        dv.promptTitle, dv.prompt = msg
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(intervalo)


def nome(wb, nome_, ref):
    wb.defined_names[nome_] = DefinedName(nome_, attr_text=ref)


def faixa(aba, col, n):
    return f"{aba}!${col}${L0}:${col}${L0 + n - 1}"


wb = Workbook()

# ---------------------------------------------------------------- abas (ordem de uso no dia a dia)
inicio = wb.active
inicio.title = "Início"
resumo = wb.create_sheet("Resumo")
vendas = wb.create_sheet("Vendas")
recebi = wb.create_sheet("Recebi")
clientes = wb.create_sheet("Clientes")
produtos = wb.create_sheet("Produtos")
receitas = wb.create_sheet("Receitas")
ingred = wb.create_sheet("Ingredientes")
mes = wb.create_sheet("Meu Mês")

# Faixas usadas em várias fórmulas
ING_ITEM, ING_UNID, ING_QTD, ING_PRECO = (faixa("Ingredientes", c, N_ING) for c in "BDEF")
REC_PROD, REC_CUSTO = faixa("Receitas", "B", N_REC), faixa("Receitas", "F", N_REC)
PROD_NOME, PROD_CUSTO, PROD_PRECO, PROD_SOBRA = (faixa("Produtos", c, N_PROD) for c in "BFKO")
CLI_NOME, CLI_DEVE = faixa("Clientes", "B", N_CLI), faixa("Clientes", "F", N_CLI)
VEN_DATA, VEN_CLI, VEN_PROD, VEN_QTD, VEN_TOTAL, VEN_CUSTO, VEN_SOBRA, VEN_NAHORA = (
    faixa("Vendas", c, N_VEN) for c in "BCDEGHIJ")
PAG_DATA, PAG_CLI, PAG_VALOR = (faixa("Recebi", c, N_PAG) for c in "BCD")

# ================================================================= MEU MÊS
topo(mes, "Meu Mês",
     "Quanto custa manter o seu negócio e quanto você precisa ganhar. Preencha uma vez e mude quando algo mudar.",
     4, CHOCOLATE)
mes.column_dimensions["B"].width = 50
mes.column_dimensions["C"].width = 16
mes.column_dimensions["D"].width = 58


def secao(ws, linha, texto, ate="D"):
    ws.merge_cells(f"B{linha}:{ate}{linha}")
    c = ws[f"B{linha}"]
    c.value = texto
    c.font = font(CHOCOLATE, 14, bold=True)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[linha].height = 28


secao(mes, 4, "1. Contas fixas do mês  (o que você paga todo mês, vendendo ou não)")
cabecalho(mes, 5, [("Conta", 50, True), ("Valor por mês", 16, True), ("Observação", 58, True)])
contas = [
    ("Gás de cozinha", 60, "A parte do botijão que vai para os doces/salgados"),
    ("Luz", 50, "Só a parte da cozinha e da geladeira, se der para separar"),
    ("Água", 20, ""),
    ("Internet e celular", 40, ""),
    ("MEI (DAS)", 82, "Se for MEI, o boleto do mês"),
    ("Transporte e entregas", 100, "Gasolina, ônibus, motoboy"),
    ("Manutenção de batedeira, forno, panelas", 30, "Guarde um pouco por mês para quando quebrar"),
    ("Material de limpeza", 25, "Detergente, esponja, pano"),
    ("Divulgação", 0, "Impressão de cardápio, anúncio"),
    ("Mensalidade da maquininha", 0, ""),
    ("Aluguel (a parte do espaço que usa)", 0, ""),
]
for i in range(15):
    r = 6 + i
    nome_c, valor, obs = contas[i] if i < len(contas) else (None, None, None)
    celula(mes, f"B{r}", nome_c, entrada=True)
    celula(mes, f"C{r}", valor, entrada=True, fmt=DINHEIRO)
    celula(mes, f"D{r}", obs, entrada=True, cor=APAGADO)
celula(mes, "B21", "Total das contas fixas", bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "C21", "=SUM(C6:C20)", fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "D21", "A planilha soma sozinha.", cor=APAGADO, fundo=VERDE_CLARO)

secao(mes, 23, "2. Você")
celula(mes, "B24", "Quanto preciso para viver por mês (meu salário)", bold=True)
celula(mes, "C24", 1500, entrada=True, fmt=DINHEIRO)
celula(mes, "D24", "O que você quer tirar para você, todo mês, para pagar as suas contas de casa.", cor=APAGADO)

secao(mes, 26, "3. Seu tempo")
celula(mes, "B27", "Dias que trabalho no mês", bold=True)
celula(mes, "C27", 22, entrada=True, alinh="center")
celula(mes, "D27", "Contando os dias de comprar, fazer e entregar.", cor=APAGADO)
celula(mes, "B28", "Horas por dia", bold=True)
celula(mes, "C28", 5, entrada=True, alinh="center")
celula(mes, "D28", "Fazendo, comprando, embalando, entregando e cobrando.", cor=APAGADO)
celula(mes, "B29", "Horas no mês")
celula(mes, "C29", "=C27*C28", alinh="center")
celula(mes, "D29", "", cor=APAGADO)
celula(mes, "B30", "Quanto vale a minha hora", bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "C30", "=IF(C29>0,(C21+C24)/C29,0)", fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "D30", "Contas fixas + seu salário, divididos pelas horas. Entra no preço de cada produto.",
       cor=APAGADO, fundo=VERDE_CLARO)

secao(mes, 32, "4. Porcentagens")
celula(mes, "B33", "Taxas sobre a venda", bold=True)
celula(mes, "C33", 0.03, entrada=True, fmt=PCT, alinh="center")
celula(mes, "D33", "Maquininha, app de entrega. Se só uma parte paga no cartão, use uma média.", cor=APAGADO)
celula(mes, "B34", "Perdas e desperdício", bold=True)
celula(mes, "C34", 0.05, entrada=True, fmt=PCT, alinh="center")
celula(mes, "D34", "O que sobra na panela, quebra, prova ou vira brinde.", cor=APAGADO)
celula(mes, "B35", "Lucro que eu quero", bold=True)
celula(mes, "C35", 0.15, entrada=True, fmt=PCT, alinh="center")
celula(mes, "D35", "Além do seu salário: para guardar, comprar equipamento e crescer.", cor=APAGADO)

secao(mes, 37, "Resultado")
celula(mes, "B38", "Preciso ganhar por mês (contas + meu salário)", bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "C38", "=C21+C24", fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "D38", "", fundo=VERDE_CLARO)
celula(mes, "B39", "Para isso, preciso vender por mês mais ou menos", bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "C39", "=MetaVendas", fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(mes, "D39", "Depende dos seus preços. Veja os detalhes no Resumo.", cor=APAGADO, fundo=VERDE_CLARO)
for r in list(range(6, 22)) + [24, 27, 28, 29, 30, 33, 34, 35, 38, 39]:
    mes.row_dimensions[r].height = 24

nome(wb, "ContasFixas", "'Meu Mês'!$C$21")
nome(wb, "Salario", "'Meu Mês'!$C$24")
nome(wb, "ValorHora", "'Meu Mês'!$C$30")
nome(wb, "Taxas", "'Meu Mês'!$C$33")
nome(wb, "Perdas", "'Meu Mês'!$C$34")
nome(wb, "LucroDesejado", "'Meu Mês'!$C$35")

# ================================================================= INGREDIENTES
topo(ingred, "Ingredientes e embalagens",
     "Tudo que você compra: o tamanho do pacote e quanto pagou. Quilo vira gramas (1 kg = 1000 g) e litro vira ml (1 L = 1000 ml).",
     10, CHOCOLATE)
cabecalho(ingred, 4, [
    ("O que é", 32, True), ("Tipo", 14, True), ("Medida\n(g, ml ou un)", 12, True),
    ("Quanto vem no pacote", 14, True), ("Quanto paguei", 14, True),
    ("Sai por", 13, False), ("", 11, False), ("Onde comprei", 18, True), ("Atualizado em", 14, True),
])
ingred.merge_cells("G4:H4")
grade(ingred, L0, N_ING, {
    "B": (True, None, None), "C": (True, None, "center"), "D": (True, None, "center"),
    "E": (True, "General", "center"), "F": (True, DINHEIRO, None), "G": (False, DINHEIRO, None),
    "H": (False, None, None), "I": (True, None, None), "J": (True, DATA, "center"),
})
hoje = dt.date(2026, 10, 30)
itens = [
    ("Leite condensado", "Ingrediente", "g", 395, 6.49, "Mercado"),
    ("Creme de leite", "Ingrediente", "g", 200, 3.29, "Mercado"),
    ("Chocolate em pó 50%", "Ingrediente", "g", 200, 9.90, "Mercado"),
    ("Chocolate meio amargo", "Ingrediente", "g", 1000, 49.90, "Atacado"),
    ("Chocolate granulado", "Ingrediente", "g", 500, 14.90, "Loja de festa"),
    ("Manteiga", "Ingrediente", "g", 200, 9.90, "Mercado"),
    ("Açúcar refinado", "Ingrediente", "g", 1000, 4.79, "Mercado"),
    ("Farinha de trigo", "Ingrediente", "g", 1000, 5.49, "Mercado"),
    ("Ovos", "Ingrediente", "un", 30, 19.90, "Feira"),
    ("Leite", "Ingrediente", "ml", 1000, 5.29, "Mercado"),
    ("Óleo", "Ingrediente", "ml", 900, 7.99, "Mercado"),
    ("Fermento em pó", "Ingrediente", "g", 100, 3.99, "Mercado"),
    ("Biscoito maisena", "Ingrediente", "g", 400, 6.49, "Mercado"),
    ("Coco ralado", "Ingrediente", "g", 100, 4.99, "Mercado"),
    ("Peito de frango", "Ingrediente", "g", 1000, 21.90, "Açougue"),
    ("Requeijão cremoso", "Ingrediente", "g", 400, 10.90, "Mercado"),
    ("Cebola", "Ingrediente", "g", 1000, 5.99, "Feira"),
    ("Sal", "Ingrediente", "g", 1000, 2.99, "Mercado"),
    ("Farinha de rosca", "Ingrediente", "g", 500, 6.90, "Mercado"),
    ("Pote 250 ml com tampa", "Embalagem", "un", 50, 39.90, "Loja de embalagem"),
    ("Colher descartável", "Embalagem", "un", 50, 7.90, "Loja de embalagem"),
    ("Adesivo com meu nome", "Embalagem", "un", 100, 15.00, "Gráfica"),
    ("Forminha nº 4", "Embalagem", "un", 100, 5.90, "Loja de festa"),
    ("Saquinho de celofane", "Embalagem", "un", 100, 9.90, "Loja de embalagem"),
    ("Caixa para salgado", "Embalagem", "un", 10, 12.90, "Loja de embalagem"),
]
for i, (n_, tipo, un, qtd, preco, onde) in enumerate(itens):
    r = L0 + i
    ingred[f"B{r}"], ingred[f"C{r}"], ingred[f"D{r}"] = n_, tipo, un
    ingred[f"E{r}"], ingred[f"F{r}"], ingred[f"I{r}"], ingred[f"J{r}"] = qtd, preco, onde, hoje
for r in range(L0, L0 + N_ING):
    ingred[f"G{r}"] = f'=IF(OR(B{r}="",E{r}="",F{r}="",N(E{r})=0),"",IF(OR(D{r}="g",D{r}="ml"),F{r}/E{r}*1000,F{r}/E{r}))'
    ingred[f"H{r}"] = f'=IF(G{r}="","",IF(D{r}="g","o quilo",IF(D{r}="ml","o litro","a unidade")))'
    ingred[f"H{r}"].font = font(APAGADO)
lista(ingred, '"Ingrediente,Embalagem,Outro"', f"C{L0}:C{L0 + N_ING - 1}")
lista(ingred, '"g,ml,un"', f"D{L0}:D{L0 + N_ING - 1}",
      ("Medida", "g = gramas, ml = mililitros, un = unidades (ovos, potes, colheres)"))
ingred.freeze_panes = "C5"

# ================================================================= RECEITAS
topo(receitas, "Receitas",
     "O que vai em cada produto e quanto vai. Inclua as embalagens (potes, colheres, adesivos). Linhas do mesmo produto ficam juntas.",
     7, CHOCOLATE)
cabecalho(receitas, 4, [
    ("Produto", 30, True), ("O que vai", 30, True), ("Quanto vai", 12, True),
    ("Medida", 10, False), ("Custo", 14, False), ("Observação", 30, True),
])
grade(receitas, L0, N_REC, {
    "B": (True, None, None), "C": (True, None, None), "D": (True, "General", "center"),
    "E": (False, None, "center"), "F": (False, DINHEIRO, None), "G": (True, None, None),
})
receitas_ex = {
    "Bolo de pote de chocolate": [
        ("Farinha de trigo", 240, "massa"), ("Açúcar refinado", 300, "massa"),
        ("Chocolate em pó 50%", 100, "massa e recheio"), ("Ovos", 4, "massa"), ("Leite", 240, "massa"),
        ("Óleo", 120, "massa"), ("Fermento em pó", 15, "massa"), ("Leite condensado", 790, "recheio (2 latas)"),
        ("Creme de leite", 200, "recheio"), ("Manteiga", 30, "recheio"), ("Chocolate granulado", 50, "cobertura"),
        ("Pote 250 ml com tampa", 12, ""), ("Colher descartável", 12, ""), ("Adesivo com meu nome", 12, ""),
    ],
    "Palha italiana": [
        ("Leite condensado", 790, ""), ("Chocolate em pó 50%", 60, ""), ("Manteiga", 40, ""),
        ("Biscoito maisena", 300, ""), ("Açúcar refinado", 50, "para passar por fora"),
        ("Saquinho de celofane", 20, ""), ("Adesivo com meu nome", 20, ""),
    ],
    "Brigadeiro": [
        ("Leite condensado", 395, ""), ("Chocolate em pó 50%", 40, ""), ("Manteiga", 20, ""),
        ("Chocolate granulado", 80, ""), ("Forminha nº 4", 30, ""),
    ],
    "Coxinha de frango": [
        ("Peito de frango", 800, "recheio"), ("Requeijão cremoso", 200, "recheio"), ("Cebola", 100, "recheio"),
        ("Farinha de trigo", 500, "massa"), ("Leite", 500, "massa"), ("Manteiga", 30, "massa"),
        ("Sal", 15, ""), ("Ovos", 2, "empanar"), ("Farinha de rosca", 300, "empanar"),
        ("Óleo", 450, "fritura"), ("Caixa para salgado", 2, ""),
    ],
}
r = L0
for prod, linhas in receitas_ex.items():
    for item, qtd, obs in linhas:
        receitas[f"B{r}"], receitas[f"C{r}"], receitas[f"D{r}"] = prod, item, qtd
        if obs:
            receitas[f"G{r}"] = obs
        r += 1
for r in range(L0, L0 + N_REC):
    m = f"MATCH(C{r},{ING_ITEM},0)"
    receitas[f"E{r}"] = f'=IF(C{r}="","",IFERROR(INDEX({ING_UNID},{m}),"?"))'
    receitas[f"F{r}"] = (f'=IF(OR(C{r}="",D{r}=""),"",'
                         f'IFERROR(D{r}*INDEX({ING_PRECO},{m})/INDEX({ING_QTD},{m}),"Cadastre em Ingredientes"))')
lista(receitas, f"={PROD_NOME}", f"B{L0}:B{L0 + N_REC - 1}",
      ("Produto", "Escolha da lista. Para um produto novo, cadastre antes na aba Produtos."))
lista(receitas, f"={ING_ITEM}", f"C{L0}:C{L0 + N_REC - 1}",
      ("O que vai", "Escolha da lista. Se não estiver, cadastre na aba Ingredientes."))
receitas.conditional_formatting.add(
    f"F{L0}:F{L0 + N_REC - 1}",
    FormulaRule(formula=[f'ISTEXT(F{L0})*(F{L0}<>"")'], font=Font(color=LARANJA, bold=True), fill=fill(LARANJA_CLARO)))
receitas.freeze_panes = "C5"

# ================================================================= PRODUTOS
topo(produtos, "Produtos e preços",
     "Quanto custa cada unidade e quanto você deveria cobrar. Preencha as colunas rosinha; o resto a planilha calcula.",
     14, CHOCOLATE)
cabecalho(produtos, 4, [
    ("Produto", 28, True), ("A receita rende\n(unidades)", 13, True), ("Tempo para fazer a receita (horas)", 14, True),
    ("Ingredientes e embalagem da receita", 15, False), ("Custo por unidade\n(ingredientes, embalagem, perdas)", 17, False),
    ("Seu tempo e contas por unidade", 15, False), ("Custo completo por unidade", 15, False),
    ("Preço mínimo\n(sem prejuízo)", 14, False), ("Preço sugerido", 14, False), ("Preço que eu cobro", 14, True),
    ("Lucro por unidade", 14, False), ("Margem de lucro", 12, False), ("Como está", 22, False),
])
grade(produtos, L0, N_PROD, {
    "B": (True, None, None), "C": (True, "General", "center"), "D": (True, "General", "center"),
    "E": (False, DINHEIRO, None), "F": (False, DINHEIRO, None), "G": (False, DINHEIRO, None),
    "H": (False, DINHEIRO, None), "I": (False, DINHEIRO, None), "J": (False, DINHEIRO, None),
    "K": (True, DINHEIRO, None), "L": (False, DINHEIRO, None), "M": (False, PCT, "center"),
    "N": (False, None, "center"),
})
for col in "HJ":
    for r in range(L0, L0 + N_PROD):
        produtos[f"{col}{r}"].font = font(TINTA, bold=True)
prods_ex = [("Bolo de pote de chocolate", 12, 3, 12), ("Palha italiana", 20, 1.5, 5),
            ("Brigadeiro", 30, 1.5, 2.5), ("Coxinha de frango", 20, 3, 5)]
for i, (n_, rende, horas, preco) in enumerate(prods_ex):
    r = L0 + i
    produtos[f"B{r}"], produtos[f"C{r}"], produtos[f"D{r}"], produtos[f"K{r}"] = n_, rende, horas, preco
for r in range(L0, L0 + N_PROD):
    ok = f'AND(B{r}<>"",N(C{r})>0)'
    produtos[f"E{r}"] = f'=IF(B{r}="","",SUMIFS({REC_CUSTO},{REC_PROD},B{r}))'
    produtos[f"F{r}"] = f'=IF({ok},E{r}*(1+Perdas)/C{r},"")'
    produtos[f"G{r}"] = f'=IF({ok},N(D{r})*ValorHora/C{r},"")'
    produtos[f"H{r}"] = f'=IF({ok},F{r}+G{r},"")'
    produtos[f"I{r}"] = f'=IF(H{r}="","",H{r}/(1-Taxas))'
    sug = f"H{r}/(1-Taxas-LucroDesejado)"
    # Arredonda para cima: de 10 em 10 centavos até R$ 3, depois de 50 em 50
    produtos[f"J{r}"] = (f'=IF(H{r}="","",IF(Taxas+LucroDesejado>=1,"",'
                         f'IF({sug}<3,CEILING({sug},0.1),CEILING({sug},0.5))))')
    produtos[f"L{r}"] = f'=IF(OR(H{r}="",N(K{r})=0),"",K{r}-H{r}-K{r}*Taxas)'
    produtos[f"M{r}"] = f'=IF(L{r}="","",L{r}/K{r})'
    produtos[f"N{r}"] = (f'=IF(B{r}="","",IF(N(C{r})=0,"Falta o rendimento",IF(N(E{r})=0,"Falta a receita",'
                         f'IF(N(K{r})=0,"Falta o seu preço",IF(K{r}<F{r}/(1-Taxas),"Prejuízo",'
                         f'IF(K{r}<I{r},"Não paga o seu tempo",IF(K{r}<J{r}-0.005,"Abaixo do sugerido","Preço bom")))))))')
    # Coluna escondida: de cada R$ 1 vendido, quanto sobra para contas e salário (usada na meta do Resumo)
    produtos[f"O{r}"] = f'=IF(OR(F{r}="",N(K{r})=0),"",(K{r}-F{r})/K{r}-Taxas)'
produtos.column_dimensions["O"].hidden = True
faixa_n = f"N{L0}:N{L0 + N_PROD - 1}"
for texto, cor_f, cor_t in [("Preço bom", VERDE_CLARO, VERDE), ("Prejuízo", LARANJA_CLARO, LARANJA),
                            ("Não paga o seu tempo", LARANJA_CLARO, LARANJA),
                            ("Abaixo do sugerido", "FFF3D6", "8A5A00")]:
    produtos.conditional_formatting.add(
        faixa_n, FormulaRule(formula=[f'N{L0}="{texto}"'], fill=fill(cor_f), font=Font(color=cor_t, bold=True)))
produtos.conditional_formatting.add(
    f"L{L0}:M{L0 + N_PROD - 1}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color=LARANJA, bold=True)))
produtos.freeze_panes = "C5"

# ================================================================= CLIENTES
topo(clientes, "Clientes",
     "Quem compra de você. O quanto cada um pegou, pagou e ainda deve aparece sozinho, pelas abas Vendas e Recebi.",
     8, ROSA)
cabecalho(clientes, 4, [
    ("Nome", 28, True), ("WhatsApp", 18, True), ("Já pegou", 14, False), ("Já pagou", 14, False),
    ("Está devendo", 15, False), ("Última compra", 14, False), ("Situação", 16, False),
])
grade(clientes, L0, N_CLI, {
    "B": (True, None, None), "C": (True, "@", None), "D": (False, DINHEIRO, None), "E": (False, DINHEIRO, None),
    "F": (False, DINHEIRO, None), "G": (False, DATA, "center"), "H": (False, None, "center"),
})
for i, (n_, tel) in enumerate([("Ana (RH)", "(11) 90000-0001"), ("Bruno da portaria", "(11) 90000-0002"),
                               ("Carla", "(11) 90000-0003"), ("Dona Lúcia", "(11) 90000-0004"),
                               ("João vizinho", "(11) 90000-0005")]):
    clientes[f"B{L0 + i}"], clientes[f"C{L0 + i}"] = n_, tel
for r in range(L0, L0 + N_CLI):
    clientes[f"D{r}"] = f'=IF(B{r}="","",SUMIFS({VEN_TOTAL},{VEN_CLI},B{r}))'
    clientes[f"E{r}"] = (f'=IF(B{r}="","",SUMIFS({PAG_VALOR},{PAG_CLI},B{r})'
                         f'+SUMIFS({VEN_TOTAL},{VEN_CLI},B{r},{VEN_NAHORA},"Sim"))')
    clientes[f"F{r}"] = f'=IF(B{r}="","",D{r}-E{r})'
    clientes[f"G{r}"] = f'=IF(B{r}="","",IF(COUNTIF({VEN_CLI},B{r})=0,"",_xlfn.MAXIFS({VEN_DATA},{VEN_CLI},B{r})))'
    clientes[f"H{r}"] = (f'=IF(B{r}="","",IF(F{r}>0.005,"Deve",IF(F{r}<-0.005,"Tem crédito",'
                         f'IF(D{r}>0,"Em dia","Ainda não comprou"))))')
    clientes[f"F{r}"].font = font(TINTA, bold=True)
faixa_h = f"H{L0}:H{L0 + N_CLI - 1}"
clientes.conditional_formatting.add(
    faixa_h, FormulaRule(formula=[f'H{L0}="Deve"'], fill=fill(LARANJA_CLARO), font=Font(color=LARANJA, bold=True)))
clientes.conditional_formatting.add(
    faixa_h, FormulaRule(formula=[f'H{L0}="Em dia"'], fill=fill(VERDE_CLARO), font=Font(color=VERDE, bold=True)))
clientes.conditional_formatting.add(
    f"F{L0}:F{L0 + N_CLI - 1}", CellIsRule(operator="greaterThan", formula=["0.005"], font=Font(color=LARANJA, bold=True)))
clientes.freeze_panes = "C5"

# ================================================================= VENDAS
topo(vendas, "Vendas",
     "Quem pegou o quê. Uma linha para cada produto. Se deixar o preço em branco, usa o preço da aba Produtos.",
     10, ROSA)
cabecalho(vendas, 4, [
    ("Data", 13, True), ("Cliente", 24, True), ("Produto", 28, True), ("Quantas", 10, True),
    ("Preço da unidade", 13, True), ("Total", 13, False), ("Custo para fazer", 13, False),
    ("Sobrou", 13, False), ("Pagou na hora?", 11, True),
])
grade(vendas, L0, N_VEN, {
    "B": (True, DATA, "center"), "C": (True, None, None), "D": (True, None, None), "E": (True, "General", "center"),
    "F": (True, DINHEIRO, None), "G": (False, DINHEIRO, None), "H": (False, DINHEIRO, None),
    "I": (False, DINHEIRO, None), "J": (True, None, "center"),
})
vendas_ex = [
    (dt.date(2026, 10, 2), "Ana (RH)", "Bolo de pote de chocolate", 2, 12, ""),
    (dt.date(2026, 10, 2), "Bruno da portaria", "Brigadeiro", 6, 2.5, "Sim"),
    (dt.date(2026, 10, 6), "Carla", "Palha italiana", 3, 5, ""),
    (dt.date(2026, 10, 9), "Dona Lúcia", "Coxinha de frango", 10, 5, ""),
    (dt.date(2026, 10, 9), "Dona Lúcia", "Brigadeiro", 10, 2.5, ""),
    (dt.date(2026, 10, 14), "Ana (RH)", "Palha italiana", 2, 5, ""),
    (dt.date(2026, 10, 16), "João vizinho", "Bolo de pote de chocolate", 1, 12, "Sim"),
    (dt.date(2026, 10, 21), "Carla", "Bolo de pote de chocolate", 2, 12, ""),
    (dt.date(2026, 10, 23), "Bruno da portaria", "Coxinha de frango", 4, 5, ""),
    (dt.date(2026, 10, 28), "Ana (RH)", "Brigadeiro", 4, 2.5, ""),
]
for i, (d, c, p, q, pr, nh) in enumerate(vendas_ex):
    r = L0 + i
    vendas[f"B{r}"], vendas[f"C{r}"], vendas[f"D{r}"], vendas[f"E{r}"], vendas[f"F{r}"] = d, c, p, q, pr
    if nh:
        vendas[f"J{r}"] = nh
for r in range(L0, L0 + N_VEN):
    m = f"MATCH(D{r},{PROD_NOME},0)"
    vendas[f"G{r}"] = (f'=IF(OR(C{r}="",D{r}="",N(E{r})=0),"",'
                       f'E{r}*IF(F{r}="",IFERROR(N(INDEX({PROD_PRECO},{m})),0),F{r}))')
    vendas[f"H{r}"] = f'=IF(G{r}="","",E{r}*IFERROR(N(INDEX({PROD_CUSTO},{m})),0))'
    vendas[f"I{r}"] = f'=IF(G{r}="","",G{r}-H{r})'
lista(vendas, f"={CLI_NOME}", f"C{L0}:C{L0 + N_VEN - 1}", ("Cliente", "Escolha da lista ou cadastre na aba Clientes."))
lista(vendas, f"={PROD_NOME}", f"D{L0}:D{L0 + N_VEN - 1}", ("Produto", "Escolha da lista da aba Produtos."))
lista(vendas, '"Sim,Não"', f"J{L0}:J{L0 + N_VEN - 1}")
vendas.freeze_panes = "B5"

# ================================================================= RECEBI
topo(recebi, "Recebi",
     "Cada vez que alguém te pagar, uma linha aqui. Pode ser tudo ou só uma parte do que a pessoa deve.",
     6, ROSA)
cabecalho(recebi, 4, [("Data", 13, True), ("Cliente", 24, True), ("Quanto", 14, True),
                      ("Como pagou", 14, True), ("Observação", 32, True)])
grade(recebi, L0, N_PAG, {
    "B": (True, DATA, "center"), "C": (True, None, None), "D": (True, DINHEIRO, None),
    "E": (True, None, "center"), "F": (True, None, None),
})
for i, (d, c, v, f_, o) in enumerate([
        (dt.date(2026, 10, 10), "Ana (RH)", 24, "Pix", ""),
        (dt.date(2026, 10, 15), "Carla", 15, "Dinheiro", ""),
        (dt.date(2026, 10, 20), "Dona Lúcia", 50, "Pix", "Pagou uma parte"),
        (dt.date(2026, 10, 25), "Bruno da portaria", 20, "Cartão", "")]):
    r = L0 + i
    recebi[f"B{r}"], recebi[f"C{r}"], recebi[f"D{r}"], recebi[f"E{r}"] = d, c, v, f_
    if o:
        recebi[f"F{r}"] = o
lista(recebi, f"={CLI_NOME}", f"C{L0}:C{L0 + N_PAG - 1}", ("Cliente", "Escolha da lista da aba Clientes."))
lista(recebi, '"Pix,Dinheiro,Cartão,Na hora"', f"E{L0}:E{L0 + N_PAG - 1}")
recebi.freeze_panes = "B5"

# ================================================================= RESUMO
topo(resumo, "Resumo",
     "Quanto você vendeu, quanto já entrou, quem está devendo e se o mês pagou as suas contas e o seu salário.",
     5, ROSA)
for col, larg in zip("BCDE", (46, 17, 17, 50)):
    resumo.column_dimensions[col].width = larg

celula(resumo, "B4", "Mês que quero ver", bold=True, size=12)
celula(resumo, "C4", "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)", entrada=True, fmt="mm/yyyy", alinh="center", bold=True, size=12)
celula(resumo, "D4", "Começa no mês de hoje. Para ver outro, escreva a data (ex.: 01/09/2026).", cor=APAGADO)
resumo.merge_cells("D4:E4")
resumo.row_dimensions[4].height = 30
no_mes = f'{VEN_DATA},">="&$C$4,{VEN_DATA},"<="&EOMONTH($C$4,0)'
pag_mes = f'{PAG_DATA},">="&$C$4,{PAG_DATA},"<="&EOMONTH($C$4,0)'


def linha_resumo(r, rotulo, f_mes, f_tudo, explica, destaque=False, fmt=DINHEIRO):
    fundo = VERDE_CLARO if destaque else None
    cor = VERDE if destaque else TINTA
    celula(resumo, f"B{r}", rotulo, bold=True, fundo=fundo, cor=cor)
    celula(resumo, f"C{r}", f_mes, fmt=fmt, bold=destaque, fundo=fundo, cor=cor)
    celula(resumo, f"D{r}", f_tudo, fmt=fmt, bold=destaque, fundo=fundo, cor=cor)
    celula(resumo, f"E{r}", explica, cor=APAGADO, fundo=fundo)
    resumo.row_dimensions[r].height = 26


secao(resumo, 6, "Dinheiro na rua", "E")
celula(resumo, "B7", "Falta receber de todo mundo", bold=True, fundo=LARANJA_CLARO, cor=LARANJA, size=13)
celula(resumo, "C7", f'=SUMIF({CLI_DEVE},">0")', fmt=DINHEIRO, bold=True, fundo=LARANJA_CLARO, cor=LARANJA, size=13)
celula(resumo, "D7", f'=COUNTIF({CLI_DEVE},">0.005")&" devendo"', fundo=LARANJA_CLARO, cor=LARANJA, bold=True)
celula(resumo, "E7", "Veja quem na aba Clientes.", cor=APAGADO, fundo=LARANJA_CLARO)
resumo.row_dimensions[7].height = 30

secao(resumo, 9, "Vendas", "E")
for col, t in zip("BCDE", ("", "Neste mês", "Desde o começo", "")):
    c = celula(resumo, f"{col}10", t, bold=True, cor=BRANCO, fundo=CHOCOLATE, alinh="center")
linha_resumo(11, "Vendi", f"=SUMIFS({VEN_TOTAL},{no_mes})", f"=SUM({VEN_TOTAL})",
             "Tudo que as pessoas pegaram, pago ou fiado.")
linha_resumo(12, "Já recebi", f'=SUMIFS({PAG_VALOR},{pag_mes})+SUMIFS({VEN_TOTAL},{VEN_NAHORA},"Sim",{no_mes})',
             f'=SUM({PAG_VALOR})+SUMIFS({VEN_TOTAL},{VEN_NAHORA},"Sim")', "O dinheiro que já entrou.")
linha_resumo(13, "Gastei com ingredientes e embalagens", f"=SUMIFS({VEN_CUSTO},{no_mes})", f"=SUM({VEN_CUSTO})",
             "Pelo custo de cada produto (aba Produtos).")
linha_resumo(14, "Taxas (maquininha, app)", "=C11*Taxas", "=D11*Taxas", "Uma estimativa, pela % da aba Meu Mês.")
linha_resumo(15, "Sobrou das vendas", "=C11-C13-C14", "=D11-D13-D14",
             "O que fica para pagar as contas e o seu salário.", destaque=True)

secao(resumo, 17, "O mês pagou as contas e o seu salário?", "E")
celula(resumo, "B18", "Contas fixas do mês", bold=True)
celula(resumo, "C18", "=ContasFixas", fmt=DINHEIRO)
celula(resumo, "D18", "", fundo=BRANCO)
celula(resumo, "E18", "Da aba Meu Mês.", cor=APAGADO)
celula(resumo, "B19", "Ficou para mim", bold=True, fundo=VERDE_CLARO, cor=VERDE, size=13)
celula(resumo, "C19", "=C15-C18", fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE, size=13)
celula(resumo, "D19", "", fundo=VERDE_CLARO)
celula(resumo, "E19", "Sobrou das vendas menos as contas fixas (contando o fiado que ainda vai entrar).",
       cor=APAGADO, fundo=VERDE_CLARO)
celula(resumo, "B20", "Meu salário", bold=True)
celula(resumo, "C20", "=Salario", fmt=DINHEIRO)
celula(resumo, "D20", "", fundo=BRANCO)
celula(resumo, "E20", "O que você quer tirar por mês (aba Meu Mês).", cor=APAGADO)
resumo.merge_cells("B21:E21")
celula(resumo, "B21", ('=IF(C11=0,"Ainda não tem vendas neste mês.",IF(C19>=C20,'
                       '"Bateu o seu salário! E ainda sobrou "&"R$ "&FIXED(C19-C20,2)&" de lucro.",'
                       'IF(C19>0,"Faltam "&"R$ "&FIXED(C20-C19,2)&" para chegar no seu salário.",'
                       '"As vendas ainda não pagaram as contas fixas do mês.")))'),
       bold=True, size=13, cor=CHOCOLATE, fundo=ROSA_FUNDO, alinh="center")
for col in "CDE":
    resumo[f"{col}21"].fill = fill(ROSA_FUNDO)
    resumo[f"{col}21"].border = BORDA
for r in (18, 19, 20):
    resumo.row_dimensions[r].height = 26
resumo.row_dimensions[21].height = 34

secao(resumo, 23, "Meta do mês", "E")
# De cada R$ 1 vendido, quanto sobra para contas e salário: do mês, ou a média dos produtos se não vendeu nada
celula(resumo, "B24", "De cada R$ 1,00 que vendo, sobra", bold=True)
celula(resumo, "C24", f'=IF(C11>0,(C11-C13)/C11-Taxas,IFERROR(AVERAGE({PROD_SOBRA}),0))', fmt=PCT, alinh="center")
celula(resumo, "D24", "", fundo=BRANCO)
celula(resumo, "E24", "Para pagar contas e salário. Quanto maior, melhor.", cor=APAGADO)
celula(resumo, "B25", "Para pagar contas + salário, preciso vender", bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(resumo, "C25", '=IF(C24>0,(ContasFixas+Salario)/C24,"")', fmt=DINHEIRO, bold=True, fundo=VERDE_CLARO, cor=VERDE)
celula(resumo, "D25", "", fundo=VERDE_CLARO)
celula(resumo, "E25", "Por mês.", cor=APAGADO, fundo=VERDE_CLARO)
celula(resumo, "B26", '="Para ter também "&TEXT(LucroDesejado,"0%")&" de lucro, preciso vender"', bold=True)
celula(resumo, "C26", '=IF(C24-LucroDesejado>0,(ContasFixas+Salario)/(C24-LucroDesejado),"")', fmt=DINHEIRO, bold=True)
celula(resumo, "D26", "", fundo=BRANCO)
celula(resumo, "E26", '=IF(C26="","Com os preços de hoje não dá esse lucro. Veja a aba Produtos.","Por mês.")', cor=APAGADO)
celula(resumo, "B27", "Já vendi da meta", bold=True)
celula(resumo, "C27", '=IF(N(C25)=0,"",C11/C25)', fmt=PCT, alinh="center", bold=True)
celula(resumo, "D27", '=IF(C27="","",REPT("■",MIN(10,ROUND(C27*10,0)))&REPT("□",10-MIN(10,ROUND(C27*10,0))))',
       cor=ROSA, alinh="center", size=13)
celula(resumo, "E27", '=IF(C27="","",IF(C11>=C25,"Meta batida!","Faltam "&"R$ "&FIXED(C25-C11,2)&" em vendas."))',
       cor=APAGADO)
for r in range(24, 28):
    resumo.row_dimensions[r].height = 26
nome(wb, "MetaVendas", "Resumo!$C$25")

secao(resumo, 29, "Por produto, neste mês", "E")
for col, t in zip("BCDE", ("Produto", "Unidades", "Vendi", "Sobrou")):
    celula(resumo, f"{col}30", t, bold=True, cor=BRANCO, fundo=CHOCOLATE, alinh="center")
for i in range(N_PROD):
    r, rp = 31 + i, L0 + i
    celula(resumo, f"B{r}", f'=IF(Produtos!B{rp}="","",Produtos!B{rp})')
    celula(resumo, f"C{r}", f'=IF(B{r}="","",SUMIFS({VEN_QTD},{VEN_PROD},B{r},{no_mes}))', fmt="General", alinh="center")
    celula(resumo, f"D{r}", f'=IF(B{r}="","",SUMIFS({VEN_TOTAL},{VEN_PROD},B{r},{no_mes}))', fmt=DINHEIRO)
    celula(resumo, f"E{r}", f'=IF(B{r}="","",SUMIFS({VEN_SOBRA},{VEN_PROD},B{r},{no_mes}))', fmt=DINHEIRO)
resumo.freeze_panes = "A5"

# ================================================================= INÍCIO
topo(inicio, "Caderninho dos Doces",
     "Planilha de custos, preços e fiado. Serve para doces, salgados, marmitas, bolos, artesanato: tudo que você vende e recebe depois.",
     4, ROSA)
inicio.column_dimensions["B"].width = 6
inicio.column_dimensions["C"].width = 24
inicio.column_dimensions["D"].width = 86
secao(inicio, 4, "Na primeira vez, preencha nesta ordem:")
passos = [
    ("Meu Mês", "Suas contas fixas, quanto você precisa para viver e quantas horas trabalha. A planilha calcula quanto vale a sua hora."),
    ("Ingredientes", "Tudo que você compra: ingredientes e embalagens, com o tamanho do pacote e o preço que pagou."),
    ("Produtos", "O nome de cada coisa que você vende, quanto a receita rende e quanto tempo leva."),
    ("Receitas", "O que vai em cada produto e quanto vai. O custo aparece sozinho, e o preço sugerido aparece em Produtos."),
    ("Clientes", "Quem compra de você."),
]
dia_a_dia = [
    ("Vendas", "Quem pegou o quê, quantos e por quanto."),
    ("Recebi", "Quem te pagou, quanto e como (Pix, dinheiro, cartão)."),
    ("Resumo", "Quanto vendeu, quanto já recebeu, quem deve e se o mês pagou suas contas e o seu salário."),
]


def passo(r, num, aba, texto):
    celula(inicio, f"B{r}", num, bold=True, cor=BRANCO, fundo=ROSA, alinh="center", size=13)
    celula(inicio, f"C{r}", aba, bold=True, cor=CHOCOLATE, size=12)
    celula(inicio, f"D{r}", texto)
    inicio.row_dimensions[r].height = 34


for i, (aba, txt) in enumerate(passos):
    passo(5 + i, i + 1, aba, txt)
secao(inicio, 11, "No dia a dia:")
for i, (aba, txt) in enumerate(dia_a_dia):
    passo(12 + i, "•", aba, txt)

secao(inicio, 16, "As cores")
for r, (cor_f, cor_t, t1, t2) in enumerate([
        (ROSA_CLARO, TINTA, "Rosinha", "Você preenche. Pode escrever por cima à vontade."),
        (BRANCO, TINTA, "Branco", "A planilha calcula sozinha. Não precisa mexer."),
        (VERDE_CLARO, VERDE, "Verde", "Os números mais importantes."),
        (LARANJA_CLARO, LARANJA, "Laranja", "Atenção: alguém devendo ou um preço que dá prejuízo.")], start=17):
    celula(inicio, f"B{r}", "", fundo=cor_f)
    celula(inicio, f"C{r}", t1, bold=True, cor=cor_t, fundo=cor_f)
    celula(inicio, f"D{r}", t2)
    inicio.row_dimensions[r].height = 26

secao(inicio, 22, "Como a planilha chega no preço")
contas_preco = [
    ("1", "Custo da unidade", "Ingredientes + embalagem da receita, mais as perdas, dividido por quantas unidades rende."),
    ("2", "Seu tempo e contas", "Horas para fazer a receita × quanto vale a sua hora, dividido pelo rendimento. "
                                "É isso que paga o gás, a luz e o seu salário."),
    ("3", "Preço sugerido", "Custo completo, mais as taxas e o lucro que você quer. Arredondado para cima."),
    ("4", "Como está", "Compara com o preço que você cobra hoje: Preço bom, Abaixo do sugerido, "
                       "Não paga o seu tempo ou Prejuízo."),
]
for i, (n_, t1, t2) in enumerate(contas_preco):
    r = 23 + i
    celula(inicio, f"B{r}", n_, bold=True, cor=BRANCO, fundo=CHOCOLATE, alinh="center")
    celula(inicio, f"C{r}", t1, bold=True, cor=CHOCOLATE)
    celula(inicio, f"D{r}", t2)
    inicio.row_dimensions[r].height = 34

secao(inicio, 28, "Dicas")
dicas = [
    "As linhas que já vêm preenchidas são exemplos. Escreva por cima ou apague.",
    "Quilo vira gramas e litro vira ml: 1 kg = 1000 g, 1 litro = 1000 ml. Ovos, potes e colheres são em unidades (un).",
    "Mudou o preço de um ingrediente? Mude só na aba Ingredientes: receitas, custos e preços sugeridos se atualizam sozinhos.",
    "Na aba Vendas, escreva o preço que você cobrou. Assim, se mudar o preço depois, as vendas antigas não mudam.",
    "Em breve o app Caderninho dos Doces vai preencher esta planilha sozinho.",
]
for i, t in enumerate(dicas):
    r = 29 + i
    celula(inicio, f"B{r}", "•", bold=True, cor=ROSA, alinh="center", size=13)
    inicio.merge_cells(f"C{r}:D{r}")
    celula(inicio, f"C{r}", t)
    inicio.row_dimensions[r].height = 30

# Impressão: cabe na largura da folha, deitada
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(SAIDA)
print("ok:", SAIDA)
