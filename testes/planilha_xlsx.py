#!/usr/bin/env python3
"""Ajudante do teste da planilha (testes/planilha.test.js). Precisa de LibreOffice (soffice) e openpyxl.

  valores <arquivo.xlsx>       abre no LibreOffice, que refaz todas as fórmulas, e devolve os números em JSON
  editar <entrada> <saida>     faz na planilha as mudanças de quem mexe no Excel (o roteiro da etapa 3)
"""
import datetime
import json
import os
import subprocess
import sys
import tempfile

import openpyxl


def recalcular(arq):
    pasta = tempfile.mkdtemp()
    wb = openpyxl.load_workbook(arq)
    # O Resumo mostra o mês de hoje; fica em outubro de 2026 (o mês dos exemplos) para o teste não depender da data.
    wb['Resumo']['C4'] = datetime.datetime(2026, 10, 1)
    entrada = os.path.join(pasta, 'entrada.xlsx')
    wb.save(entrada)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', os.path.join(pasta, 'out'), entrada],
                   check=True, capture_output=True, timeout=170)
    return openpyxl.load_workbook(os.path.join(pasta, 'out', 'entrada.xlsx'), data_only=True)


def valores(arq):
    wb = recalcular(arq)
    P, C, R = wb['Produtos'], wb['Clientes'], wb['Resumo']
    produtos = {P[f'B{r}'].value: {c: P[f'{c}{r}'].value for c in 'GHIJKLMNO'} for r in range(5, 105) if P[f'B{r}'].value}
    clientes = {C[f'B{r}'].value: C[f'H{r}'].value for r in range(5, 305) if C[f'B{r}'].value}
    resumo = {k: R[c].value for k, c in [('falta', 'C7'), ('vendi', 'C11'), ('recebi', 'C12'), ('gastei', 'C13'), ('taxas', 'C14'),
                                         ('sobrou', 'C15'), ('contas', 'C18'), ('ficou', 'C19'), ('sobra', 'C24'), ('vender', 'C25'),
                                         ('comLucro', 'C26')]}
    por_produto = {R[f'B{r}'].value: [R[f'C{r}'].value, R[f'D{r}'].value, R[f'E{r}'].value]
                   for r in range(31, 131) if R[f'B{r}'].value and R[f'C{r}'].value}
    V = wb['Vendas']
    vendas = sum(1 for r in range(5, 3005) if V[f'C{r}'].value and (V[f'D{r}'].value or V[f'F{r}'].value))
    return {'vendas': vendas, 'produtos': produtos, 'clientes': clientes, 'resumo': resumo, 'porProduto': por_produto,
            'valorHora': wb['Meu Mês']['C30'].value}


def linha(ws, nome, col='B', de=5, ate=3004):
    for r in range(de, ate + 1):
        if ws[f'{col}{r}'].value == nome:
            return r
    raise SystemExit(f'não achei {nome!r} na aba {ws.title}')


def vazia(ws, de=5):
    r = de
    while ws[f'B{r}'].value not in (None, '') or ws[f'C{r}'].value not in (None, ''):
        r += 1
    return r


def editar(entrada, saida):
    wb = openpyxl.load_workbook(entrada)
    M, P, Rc, MM, V, Re, C = (wb[n] for n in ('Materiais', 'Produtos', 'Receitas', 'Meu Mês', 'Vendas', 'Recebi', 'Clientes'))
    M[f'F{linha(M, "Peito de frango")}'] = 25.9                     # o frango ficou mais caro
    r = vazia(M)
    for c, v in zip('BCDEFI', ['Granulado colorido', 'Ingrediente', 'g', 150, 8.9, 'Mercado']):
        M[f'{c}{r}'] = v                                              # material novo
    P[f'C{linha(P, "Coxinha de frango")}'] = 25                      # a receita de coxinha passou a render 25
    P[f'L{linha(P, "Brigadeiro")}'] = 3                               # brigadeiro passou a custar R$ 3,00
    r = vazia(Rc)
    Rc[f'B{r}'], Rc[f'C{r}'], Rc[f'D{r}'] = 'Brigadeiro', 'Granulado colorido', 30   # receita ganhou o granulado
    MM[f'C{linha(MM, "Horas por dia", ate=40)}'] = 6
    MM[f'C{linha(MM, "Taxas sobre a venda", ate=40)}'] = 0.04
    dia = datetime.datetime(2026, 10, 30)
    r = vazia(V)
    V[f'B{r}'], V[f'C{r}'], V[f'D{r}'], V[f'E{r}'] = dia, 'Bruno', 'Coxinha de frango', 3      # sem preço: usa o do produto
    V[f'B{r+1}'], V[f'C{r+1}'], V[f'D{r+1}'], V[f'E{r+1}'], V[f'J{r+1}'] = dia, 'Dona Cida', 'Brigadeiro', 2, 'Sim'
    r = vazia(Re)
    for c, v in zip('BCDEF', [dia, 'Lúcia', 20, 'Pix', 'Pagou o resto']):
        Re[f'{c}{r}'] = v
    C[f'D{linha(C, "Carla", ate=304)}'] = '3º ano C'
    wb.save(saida)
    return {'ok': True}


if __name__ == '__main__':
    acao = sys.argv[1]
    print(json.dumps(valores(sys.argv[2]) if acao == 'valores' else editar(sys.argv[2], sys.argv[3]), default=str, ensure_ascii=False))
